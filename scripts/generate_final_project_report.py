from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_PATH = PROJECT_ROOT / "reports" / "项目总报告.docx"


def set_run_font(run, size=10.5, bold=False):
    """设置 Word 中文字体。"""
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_title(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run)


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=9, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            run = cell.paragraphs[0].add_run(str(value))
            set_run_font(run, size=9)
    document.add_paragraph()


def add_figure(document, image_path, title, note):
    """插入图表、图题和简短说明。"""
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=10.5)
    title_run.italic = True

    add_body(document, note)


def add_process_flow(document):
    """用表格形式展示项目主流程。"""
    steps = [
        ("1", "原始数据", "酒店预订订单数据"),
        ("2", "数据清洗", "缺失值、异常值、重复值和时间字段处理"),
        ("3", "数据库与 SQL", "SQLite 建库、维度表和基础 SQL 分析"),
        ("4", "特征工程", "时间、预订、客户、价格、风险和收益管理特征"),
        ("5", "特征预处理", "时间窗口划分、编码、标准化和特征有效性评估"),
        ("6", "模型训练", "传统模型、深度学习模型和融合模型"),
        ("7", "模型解释", "逻辑回归系数、SHAP 和核心特征识别"),
        ("8", "业务模拟", "风险分层、干预阈值和收益测算"),
        ("9", "工程化封装", "pipeline、预测接口、API 文档和单元测试"),
    ]
    add_table(document, ["步骤", "环节", "主要内容"], steps)


def main():
    document = Document()
    add_title(document, "酒店预订行为与取消预测项目总报告")

    add_heading(document, "1. 项目背景")
    add_body(
        document,
        "酒店预订取消会影响酒店库存管理、收入预测和客户服务安排。对于 OTA 和酒店运营场景，"
        "如果能够提前识别高取消风险订单，就可以辅助制定确认提醒、担保策略、库存分配和收益管理方案。"
        "本项目围绕酒店预订订单数据，构建从数据清洗、数据库分析、特征工程到多模型预测和业务模拟的完整分析流程。",
    )
    add_process_flow(document)

    add_heading(document, "2. 项目目标")
    for item in [
        "完成原始数据质量检查、数据清洗和标准化处理。",
        "构建 SQLite 数据库和 SQL 分析脚本，支持基础业务统计。",
        "构建面向取消预测的时间、预订、客户、价格、收益管理和风险特征。",
        "训练并对比传统机器学习模型、深度学习模型和融合模型。",
        "通过模型解释识别影响预订取消的核心因素。",
        "基于预测结果设计风险分层、干预阈值模拟和业务优化建议。",
        "将核心代码封装为可复用 Python 工具包，方便团队成员复现和使用。",
    ]:
        add_bullet(document, item)

    add_heading(document, "3. 数据概况")
    add_body(
        document,
        "项目使用酒店预订订单数据，目标变量为 is_canceled，表示订单是否取消。数据字段覆盖酒店类型、"
        "提前预订天数、入住日期、入住晚数、客户人数、餐饮类型、国家或地区、市场细分、分销渠道、"
        "历史取消行为、押金类型、代理商、公司、客户类型、平均房价和预订状态等信息。"
    )
    add_body(
        document,
        "当前数据是订单级数据，没有真实 user_id。因此项目中的 RFM 分群和客户价值分析是基于客户类型、"
        "市场细分、分销渠道和是否回头客构造的客户群体近似，不代表真实单个用户画像。"
    )
    add_body(
        document,
        "清洗并构建特征后的数据规模为 110,628 行、86 列，整体取消率约为 37.08%。"
        "特征预处理后，训练集为 74,132 行，验证集为 18,628 行，测试集为 17,868 行；"
        "三份数据的取消率分别约为 36.99%、37.17% 和 37.36%，说明按时间窗口划分后正负样本比例整体保持稳定。"
    )

    add_heading(document, "4. 数据清洗与质量检查")
    for item in [
        "对原始数据进行字段完整性、缺失值、重复值和目标变量分布检查。",
        "对 children、agent、company 等字段采用 0 填补策略。",
        "对 adr 房价字段和入住天数字段进行异常值检查和处理。",
        "统一日期字段格式，构造标准日期字段，便于后续时间窗口划分和趋势分析。",
        "清洗后数据保存为 CSV 和 Parquet 格式，便于后续分析和建模读取。",
    ]:
        add_bullet(document, item)

    add_heading(document, "5. SQLite 数据库与 SQL 分析")
    add_body(
        document,
        "项目构建了 SQLite 数据库，将清洗后的数据导入预订主表，并构建客户维度表、酒店维度表和时间维度表。"
        "SQL 分析覆盖酒店预订总量与取消率、市场细分、分销渠道、月度趋势、房型与餐饮分布、国家或地区差异、"
        "提前预订天数和回头客行为等内容。"
    )

    add_heading(document, "6. 特征工程与特征预处理")
    for item in [
        "时间与季节性特征：入住月份、星期几、是否周末、节假日代理变量、淡旺季标识等。",
        "预订行为特征：总入住晚数、总客人数、是否带儿童或婴儿、房型是否匹配、是否有特殊需求、提前预订天数分组等。",
        "客户群体特征：基于客户类型、市场细分、分销渠道和是否回头客构造客户群体键，并计算历史预订次数、历史取消率和 RFM 近似价值分。",
        "价格与收益管理特征：价格与酒店、城市、客户群体历史均价的比值，价格波动指数、价格敏感度得分、预估收入、入住率代理变量和需求紧张程度。",
        "风险与业务评分特征：高风险国家或地区、高风险代理商、取消风险评分等。",
        "特征预处理：对 country、agent、hotel、customer_group_key 等高维类别字段使用平滑 Target Encoding，对低维类别字段使用 One-Hot Encoding，对数值型字段使用训练集拟合 StandardScaler。",
        "特征有效性验证：通过方差阈值、相关性分析、互信息法和随机森林重要性预评估筛选和评估特征。",
        "预处理结果：特征工程数据为 110,628 行、86 列；候选特征 82 个；编码和标准化后训练集特征 168 个，经方差阈值过滤后保留 167 个建模特征。",
    ]:
        add_bullet(document, item)

    add_heading(document, "7. 探索性数据分析")
    add_body(
        document,
        "EDA 从整体业务、取消行为、用户分群、价格、渠道和季节性多个角度展开。分析结果显示，"
        "取消风险与订单来源、押金类型、提前预订天数、历史取消行为和部分渠道特征存在明显关系。"
        "不同酒店类型、不同客户类型和不同渠道的取消率存在差异，但部分时间趋势的波动幅度较小，"
        "解释时需要结合订单量和比例尺判断。"
    )
    eda_figure_dir = PROJECT_ROOT / "reports" / "eda_figures"
    add_figure(
        document,
        eda_figure_dir / "fig02_04_hotel_monthly_trends.png",
        "图 1 两类酒店月度预订量、取消率与平均房价趋势",
        "该图从预订量、取消率和平均房价三个角度对比城市酒店和度假酒店。整体来看，城市酒店的预订量明显高于度假酒店，"
        "说明当前样本中城市酒店订单占比更高；城市酒店取消率也整体高于度假酒店，提示酒店类型本身可能与取消风险有关。"
        "两类酒店的月度取消率虽有波动，但幅度相对有限，因此解释时间趋势时需要同时结合订单量和坐标比例尺，不能只根据折线起伏判断波动很大。"
        "平均房价方面，城市酒店整体高于度假酒店，因此本项目在特征工程和建模中同时考虑酒店类型与价格水平。",
    )
    add_figure(
        document,
        eda_figure_dir / "fig09_deposit_type_cancel_rate.png",
        "图 2 押金类型与取消率",
        "该图展示不同押金类型下的订单量和取消率差异。押金类型反映订单约束条件，不同押金安排对应的取消率差异较明显，"
        "说明押金规则与取消行为之间存在较强关联。因此，押金类型在本项目中被作为重要特征，也为业务侧设计担保、押金或取消政策提供参考。",
    )
    add_figure(
        document,
        eda_figure_dir / "fig13_14_rfm_segments.png",
        "图 3 RFM 客户群体取消率与消费特征",
        "该图对比不同 RFM 近似客户群体的取消率和消费特征，用于观察客户群体价值与取消行为之间的关系。"
        "不同群体在取消率和平均消费特征上存在差异，说明客户类型、市场细分、渠道和回头客属性组合能够提供一定的行为区分度。"
        "需要注意的是，当前数据没有真实 user_id，因此这里的 RFM 是基于客户群体的近似分析，不能解释为真实单个用户画像。",
    )
    add_figure(
        document,
        eda_figure_dir / "fig16_adr_bin_cancel_rate.png",
        "图 4 房价区间与取消率",
        "该图展示不同 ADR 房价区间下的订单分布和取消率，用于观察价格水平与取消行为之间的关系。"
        "如果某些价格区间的取消率明显偏高，说明价格敏感性可能会影响用户是否保留订单；如果订单量较少的价格区间取消率波动较大，"
        "则需要谨慎解释，避免把小样本波动误判为稳定规律。该分析为价格比值、价格波动指数和价格敏感度得分等特征提供依据。",
    )
    add_figure(
        document,
        eda_figure_dir / "fig18_20_channel_analysis.png",
        "图 5 分销渠道预订量、取消率与平均房价",
        "该图从订单量、取消率和平均房价三个维度比较不同分销渠道。渠道之间的订单规模和取消率存在差异，"
        "说明订单来源会影响取消风险和收入表现。对于订单量较大且取消率偏高的渠道，需要重点关注库存占用和取消损失；"
        "对于取消率较低、收入较稳定的渠道，可以在库存紧张时优先保障。该结果直接支撑本项目的渠道库存分配和风险分层策略。",
    )
    add_figure(
        document,
        eda_figure_dir / "fig21_23_season_analysis.png",
        "图 6 淡旺季预订量、取消率与平均房价",
        "该图用于观察淡旺季对预订量、取消率和平均房价的影响。从图中看，旺季和淡季的预订量、取消率和平均房价差异都不明显，"
        "说明在当前样本和当前淡旺季划分口径下，季节性没有表现出很强的区分度。因此，淡旺季标识在本项目中更多作为辅助变量使用，"
        "不能单独作为判断取消风险或价格变化的主要依据；取消风险仍需要结合渠道、客户类型、提前预订天数和押金类型等因素综合分析。",
    )

    add_heading(document, "8. 模型构建与实验结果")
    add_body(
        document,
        "项目采用时间窗口划分训练集、验证集和测试集，避免随机划分可能带来的时间信息泄露。"
        "传统机器学习部分实现了逻辑回归、随机森林、XGBoost 和 LightGBM，并使用 Optuna 进行超参数调优。"
        "深度学习部分实现了 MLP 和 Embedding MLP。融合部分对比了加权平均融合和多组 Stacking 方案。"
    )
    add_body(
        document,
        "在样本构建阶段，项目对比了原始训练集、SMOTE 过采样、欠采样和类别权重调整四种训练口径。"
        "由于取消样本占比约 37%，类别并非极端不平衡，因此最终建模主要采用类别权重调整方式，"
        "在不改变样本时间结构的前提下缓解类别比例差异。"
    )
    add_body(
        document,
        "测试集结果显示，LightGBM 的 AUC 为 0.9485，F1 为 0.8326；XGBoost 的 AUC 为 0.9483，F1 为 0.8329，"
        "二者是当前表现最稳定的强模型。随机森林测试集 AUC 为 0.9422，逻辑回归测试集 AUC 为 0.9152，"
        "可以作为解释性较强的基线模型。深度学习模型中，Embedding MLP 的测试集 AUC 为 0.9429，高于普通 MLP 的 0.9341，"
        "但没有超过 LightGBM 和 XGBoost。"
    )
    add_body(
        document,
        "项目还按城市酒店和度假酒店拆分测试样本，分别评估模型表现。结果显示，LightGBM 和 XGBoost 在两类酒店样本上均保持较高 AUC，"
        "其中度假酒店样本的 AUC 略高，说明模型在不同酒店类型下具有一定稳定性。"
    )
    add_body(
        document,
        "融合实验中，Weighted_Average_LGBM_Heavy、Stacking_C_RF_Meta 和 Stacking_D_LGBM_Meta 的 AUC 均接近 0.9485，"
        "与 LightGBM 基本持平。其中 Stacking_D_LGBM_Meta 的 accuracy、precision 和 F1 略高，但提升幅度很小。"
        "因此当前特征体系下，融合模型可以达到强单模型水平，但相对于 LightGBM、XGBoost 的边际增益有限。"
    )
    add_table(
        document,
        ["模型类型", "主要模型", "结论"],
        [
            ["传统机器学习", "逻辑回归、随机森林、XGBoost、LightGBM", "LightGBM 和 XGBoost 整体表现较强。"],
            ["深度学习", "MLP、Embedding MLP", "可以处理结构化特征，但整体没有明显超过强树模型。"],
            ["模型融合", "加权平均、Stacking", "融合模型可接近强单模型表现，但边际提升有限。"],
        ],
    )
    add_table(
        document,
        ["模型", "AUC", "Accuracy", "Precision", "Recall", "F1", "说明"],
        [
            ["LightGBM", "0.9485", "0.8716", "0.8113", "0.8552", "0.8326", "强树模型，综合表现稳定。"],
            ["XGBoost", "0.9483", "0.8721", "0.8135", "0.8532", "0.8329", "与 LightGBM 接近，F1 略高。"],
            ["Random Forest", "0.9422", "0.8639", "0.8074", "0.8348", "0.8209", "表现稳定，但低于 boosting 模型。"],
            ["Embedding MLP", "0.9429", "0.8615", "0.7904", "0.8567", "0.8222", "优于普通 MLP，但未超过强树模型。"],
            ["MLP", "0.9341", "0.8459", "0.7648", "0.8486", "0.8045", "深度学习基线模型。"],
            ["Logistic Regression", "0.9152", "0.8196", "0.7175", "0.8532", "0.7795", "解释性较强，适合作为基线。"],
            ["Stacking_D_LGBM_Meta", "0.9485", "0.8735", "0.8193", "0.8487", "0.8337", "融合模型中综合指标略高，但提升有限。"],
        ],
    )

    add_heading(document, "9. 模型解释与核心特征")
    add_body(
        document,
        "模型解释结合逻辑回归系数和 XGBoost、LightGBM 的 SHAP 结果。综合排序靠前的特征集中在国家或地区历史取消倾向、"
        "押金类型、取消风险评分、代理商信息、停车位需求、历史取消次数、预订房型、房型是否匹配和特殊需求数量。"
        "这些特征共同反映了订单来源、历史行为、预订约束条件和服务需求对取消风险的影响。"
    )
    add_body(
        document,
        "逻辑回归系数用于判断特征对取消概率的方向性影响，例如 country_target_encoded、reserved_room_type_P、"
        "deposit_type_Non_Refund、agent_target_encoded 和 previous_cancellations_scaled 等特征系数为正，"
        "表示这些特征取值升高时模型更倾向于预测订单取消；required_car_parking_spaces_scaled 和 has_car_parking_request_scaled 等特征系数为负，"
        "表示停车需求较强的订单更倾向于被预测为不取消。"
    )
    add_body(
        document,
        "SHAP 结果用于解释 XGBoost 和 LightGBM 等树模型的特征贡献。两个模型都将押金类型、取消风险评分、国家或地区编码、"
        "代理商信息、历史取消行为、房型匹配和特殊需求数量识别为重要变量，说明这些因素不仅在线性模型中有解释力，"
        "在非线性树模型中也具有较高贡献。"
    )
    add_figure(
        document,
        PROJECT_ROOT / "reports" / "core_feature_top10.png",
        "图7 影响预订取消的 Top 10 核心特征",
        "该图综合逻辑回归系数、XGBoost SHAP 和 LightGBM SHAP 的排序结果，用于识别多个模型共同认为重要的取消风险因素。"
        "排序靠前的特征集中在国家或地区历史取消倾向、押金类型、取消风险评分、代理商信息、停车位需求、历史取消次数、"
        "房型信息和特殊需求数量，说明取消风险不是由单一变量决定，而是由订单来源、历史行为、预订约束条件和服务需求共同影响。"
    )
    add_table(
        document,
        ["核心特征", "含义", "解释"],
        [
            ["country_target_encoded", "国家或地区历史取消倾向", "反映订单来源地区与取消风险的关系。"],
            ["deposit_type_Non_Refund", "不可退押金类型", "押金规则与用户取消行为存在明显关联。"],
            ["cancellation_risk_score_scaled", "取消风险评分", "综合历史取消行为、提前预订、押金和特殊需求等信息。"],
            ["agent_target_encoded", "代理商历史取消倾向", "反映不同代理商来源订单的风险差异。"],
            ["required_car_parking_spaces_scaled", "停车位需求", "停车位需求较强的订单通常预订意图更明确。"],
            ["previous_cancellations_scaled", "历史取消次数", "历史取消行为对当前订单取消风险有解释价值。"],
            ["deposit_type_No_Deposit", "无押金类型", "无押金订单的取消约束较弱，可能影响取消风险。"],
            ["reserved_room_type_P", "预订房型 P", "反映特定预订房型与取消风险之间的关系。"],
            ["room_type_matched_scaled", "房型是否匹配", "反映预订房型和实际分配房型之间的差异。"],
            ["total_of_special_requests_scaled", "特殊需求数量", "特殊需求越多的订单通常用户意图更明确。"],
        ],
    )
    add_body(
        document,
        "预测错误样本分析进一步检查模型在哪些场景下更容易出错。结果显示，冷启动近似样本的错误率高于非冷启动近似样本，"
        "说明当历史行为信息较少时，模型对订单取消风险的判断更不稳定；特殊日期、周末和节假日前后代理场景的错误率差异相对有限，"
        "说明这些时间代理变量在当前数据中不是主要错误来源。"
    )

    add_heading(document, "10. 业务落地模拟与建议")
    add_body(
        document,
        "项目基于 LightGBM 输出取消风险分数，将订单划分为低风险、中风险、高风险和极高风险，并进一步模拟不同干预阈值下的潜在净收益。"
        "离线测算显示，当干预阈值设置为 0.30 时，可以覆盖较多真实取消订单，并在当前假设参数下获得较高模拟净收益。"
    )
    add_body(
        document,
        "在阈值为 0.30 的离线模拟中，共触达 8,839 个订单，覆盖 6,235 个实际取消订单，取消订单召回率为 93.39%，"
        "触达订单中的取消订单占比为 70.54%。在当前假设参数下，预计可减少取消损失 148,944.98，"
        "干预成本为 44,195.00，预计净收益为 104,749.98。该结果用于评估模型潜在业务价值，"
        "真实效果仍需要通过线上 A/B 测试验证。"
    )
    add_body(
        document,
        "1. 模拟 A/B 测试方案：实验对象为预测取消风险分数达到干预阈值的订单，例如离线模拟中表现较好的 0.30。"
        "在这些高风险订单中，应采用随机分组方式选取实验组和对照组，保证两组在酒店类型、渠道、入住月份、房价和提前预订天数等方面尽量接近。"
        "实验组采取确认短信、入住前二次确认、灵活改期、押金或信用卡担保等措施；对照组不采取额外干预，继续沿用原有处理流程。"
        "核心评估指标包括取消率变化、收入变化、干预成本、净收益和用户满意度相关指标。"
        "由于当前离线数据没有满意度字段，真实上线时需要额外采集投诉率、客服咨询率、短信退订率或用户评分等指标。"
    )
    add_body(
        document,
        "2. 动态定价建议：对高取消风险订单，不是简单提高价格，而是结合担保要求、取消窗口和产品选项进行调整。"
        "例如在高风险订单上增加信用卡担保或押金选项，在中风险订单上加强确认提醒，在低风险订单上保持较灵活的预订体验。"
        "这样可以在控制取消风险的同时，减少对稳定订单的额外干扰。"
    )
    add_body(
        document,
        "3. 渠道库存建议：不同渠道的取消风险和收入稳定性不同。库存紧张时，可以优先保障取消率较低、收入较稳定的渠道，"
        "同时对取消风险较高或波动较大的渠道适度收紧库存释放或提高担保要求。该策略的核心不是简单减少某类渠道订单，"
        "而是在有限库存下优先保护更确定的预订需求，降低高风险订单占用库存后取消带来的机会损失。"
    )
    add_body(
        document,
        "4. 取消政策建议：基础取消政策应保持公开透明，避免让用户感到规则不一致。在此基础上，可以通过不同产品选项实现风险分层，"
        "例如保留灵活取消房价，同时提供价格更低但取消约束更强的担保型或不可退订选项。这样既保留低风险用户的预订体验，"
        "也为高风险订单提供更清晰的约束方式。"
    )
    add_table(
        document,
        ["风险分数区间", "风险等级", "建议动作", "业务目的"],
        [
            ["0.00-0.30", "低风险", "维持常规服务，避免过度打扰。", "保护稳定订单和用户体验。"],
            ["0.30-0.50", "中风险", "发送入住提醒或确认短信。", "用低成本方式提升确认率。"],
            ["0.50-0.70", "高风险", "加强确认，并提供灵活改期选项。", "减少可挽回取消。"],
            ["0.70-1.00", "极高风险", "考虑人工确认、押金或信用卡担保。", "降低高风险订单占用库存后的机会损失。"],
        ],
    )

    add_heading(document, "11. 工程化封装")
    add_body(
        document,
        "前面阶段主要通过 Notebook 完成分析和实验，适合展示思路和结果，但如果其他人想复用其中某一步，就需要反复复制代码。"
        "因此第四阶段把常用流程整理成 Python 模块，把数据读取、模型评估、模型训练、模型融合、预测推理和业务模拟拆分到 src 目录下。"
        "这样后续只需要调用对应函数或脚本，就可以复用同一套处理逻辑。"
    )
    add_body(
        document,
        "项目提供了统一入口 `python src/pipeline.py --step feature_pipeline`，用于重新生成特征工程数据、训练集、验证集和测试集。"
        "第三阶段建模 Notebook 也已在关闭长时间训练开关的情况下完整执行通过。"
        "同时，项目使用 `python -m unittest discover -s tests -p \"test_*.py\"` 运行基础单元测试，共 7 个测试全部通过，用来检查核心函数是否可以正常调用。"
    )
    add_body(
        document,
        "预测接口放在 src/models/predict.py 中，既可以一次预测多条订单，也可以对单条订单计算取消风险。"
        "项目还补充了 docs/api_usage.md 和 docs/user_manual.md，分别说明接口怎么调用、输入输出是什么、完整流程怎么运行以及使用时需要注意什么。"
        "这样其他团队成员不需要先读完整 Notebook，也可以按照文档运行项目工具。"
    )
    add_table(
        document,
        ["模块", "路径", "用途"],
        [
            ["数据集工具", "src/models/dataset.py", "读取建模数据、拆分 X/y、统计样本分布。"],
            ["模型评估", "src/models/evaluation.py", "统一计算 AUC、accuracy、precision、recall 和 F1。"],
            ["传统模型", "src/models/traditional.py", "构建模型、复用 Optuna 参数、保存和加载模型。"],
            ["模型融合", "src/models/ensemble.py", "执行加权平均和 Stacking 融合。"],
            ["预测推理", "src/models/predict.py", "支持批量预测和单条预测。"],
            ["业务模拟", "src/business/simulation.py", "支持风险分层和干预阈值收益模拟。"],
        ],
    )

    add_heading(document, "12. 核心洞察")
    for item in [
        "取消风险与订单来源、押金类型、代理商、历史取消行为和预订约束条件关系较强。",
        "强树模型更适合当前结构化特征数据，LightGBM 和 XGBoost 在综合指标上表现较稳定。",
        "深度学习模型可以作为对比方案，但在当前数据和特征体系下没有明显超过强树模型。",
        "模型融合能够达到强单模型水平，但相对于 LightGBM、XGBoost 的提升有限。",
        "风险分层具有业务解释价值，可以用于 A/B 测试方案和收益管理策略设计。",
    ]:
        add_bullet(document, item)

    add_heading(document, "13. 项目复盘与后续优化")
    add_heading(document, "13.1 全流程复盘", level=2)
    add_body(
        document,
        "本项目从订单级原始数据出发，依次完成数据质量检查、数据清洗、SQLite 数据库构建、SQL 分析、"
        "特征工程、探索性数据分析、特征预处理、模型训练、模型融合、模型解释、业务落地模拟和工程化封装。"
        "整体流程体现了从业务问题定义到数据处理、模型验证和业务应用设计的完整数据分析路径。"
    )
    add_body(
        document,
        "在建模阶段，项目先通过时间窗口划分训练集、验证集和测试集，再使用传统机器学习、深度学习和融合模型进行对比。"
        "在业务阶段，项目没有只停留在模型指标，而是进一步将预测概率转化为风险分层、干预阈值和收益测算，"
        "使模型结果能够对应到实际运营动作。"
    )

    add_heading(document, "13.2 问题与解决方案", level=2)
    add_body(
        document,
        "本节总结项目实施过程中已经遇到、并在当前版本中采取了处理方式的问题。重点是说明本项目如何在现有数据条件下继续完成分析和建模。"
    )
    for item in [
        "原始数据没有真实 user_id，因此当前版本没有把 RFM 结果解释为真实单个用户画像，而是采用客户类型、市场细分、分销渠道和是否回头客构造客户群体近似。",
        "原始数据没有明确订单唯一标识，因此当前版本不编造订单 ID，去重检查主要基于完全重复行。",
        "country、agent 等类别字段维度较高，直接 One-Hot 会造成特征过多，因此采用平滑 Target Encoding 处理高维类别字段。",
        "随机划分可能引入时间信息泄露，因此最终采用按 arrival_date 的时间窗口划分训练集、验证集和测试集。",
        "单一 Stacking 结果不足以说明融合模型整体效果，因此补充加权平均融合和多组 Stacking 方案进行对比。",
        "业务收益无法仅通过离线数据直接证明，因此当前版本采用离线收益模拟估算潜在价值。",
    ]:
        add_bullet(document, item)

    add_heading(document, "13.3 局限性与后续优化方向", level=2)
    add_body(
        document,
        "本节总结当前版本仍然无法完全解决的问题，以及如果后续继续推进项目，可以优先补充的数据、模型和业务验证方向。"
    )
    for item in [
        "引入更多外部数据，例如天气数据、真实法定节假日数据、当地活动数据和竞争对手价格数据，增强对需求波动的解释能力。",
        "接入真实用户 ID 后，可以构建更准确的用户画像、复购行为、RFM 分群和历史取消行为特征。",
        "探索更先进的模型架构，例如 Transformer、图神经网络或序列模型，用于捕捉时间序列变化、渠道关系和用户行为序列。",
        "增加实时特征，例如实时库存、实时价格、近期渠道流量和订单确认状态，实现更接近业务现场的实时取消预测。",
        "通过线上 A/B 测试验证不同干预策略对取消率、收入和用户体验的真实影响。",
    ]:
        add_bullet(document, item)

    add_heading(document, "13.4 可复用分析框架", level=2)
    for item in [
        "先从订单级数据质量检查入手，明确目标变量、缺失值、异常值、重复值和时间字段。",
        "围绕酒店场景构建时间、渠道、客户、价格、库存需求和历史行为特征。",
        "建模前采用时间窗口划分，减少时间信息泄露，并保证验证集和测试集更接近真实预测场景。",
        "模型评估同时关注 AUC、precision、recall 和 F1，避免只看单一指标。",
        "模型解释结合统计结果、SHAP 或系数分析，并回到业务语义中解释特征影响。",
        "业务落地时先做风险分层和离线模拟，再通过 A/B 测试验证干预策略。",
    ]:
        add_bullet(document, item)

    add_heading(document, "14. 注意事项")
    for item in [
        "当前数据没有真实 user_id，用户分群相关结论不能解释为真实单个用户画像。",
        "节假日变量是基于周末和相邻日期构造的代理变量，不等同于真实法定节假日。",
        "收益测算属于离线模拟，真实效果需要通过线上 A/B 测试验证。",
        "SHAP 和特征重要性结果反映模型预测依据，不代表严格因果关系。",
        "单条预测接口要求输入字段与训练模型使用的特征字段保持一致。",
    ]:
        add_bullet(document, item)

    add_heading(document, "15. 项目产出")
    for item in [
        "阶段报告：阶段一实验报告、第二阶段实验报告、第三阶段实验报告。",
        "项目总报告：reports/项目总报告.docx。",
        "用户手册：docs/user_manual.md。",
        "API 文档：docs/api_usage.md。",
        "统一流程入口：src/pipeline.py。",
        "Notebook：数据清洗、特征工程、EDA、特征预处理和建模 Notebook。",
        "脚本模块：src/data、src/database、src/features、src/models、src/business。",
        "结果文件：SQL 结果、特征字典、模型指标、模型解释、错误样本分析和业务模拟结果。",
        "单元测试：tests/test_model_package.py。",
    ]:
        add_bullet(document, item)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
